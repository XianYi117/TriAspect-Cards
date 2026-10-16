"""把卡表 Excel 轉成 data/cards.json。用法：python3 tools/xlsx_to_json.py data/卡牌圖鑑_v1.0.2.xlsx
需要 openpyxl（pip install openpyxl）。工作表名稱：「卡牌圖鑑_v1.0.2」與「主角圖鑑」。"""
import json, re, sys, pathlib
import openpyxl
root = pathlib.Path(__file__).resolve().parent.parent
wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
card_sheet = next(ws for ws in wb if ws.title.startswith('卡牌圖鑑'))
def norm(e):
    if not e: return None
    e = str(e).strip().replace('無效對手效果', '無效化對手效果').replace('（', '(').replace('）', ')').replace('抽一張', '抽1張')
    return re.sub(r'(本次我方單位受到傷害|對手單位傷害|本次敵方傷害|本次敵方效果傷害)-(\d)', r'敵方單位傷害-\2', e)
cards = []
for r in card_sheet.iter_rows(min_row=2, values_only=True):
    if not r[0]: continue
    cards.append({'id': r[0], 'name': r[1], 'color': {'紅': 'R', '綠': 'G', '藍': 'B'}[r[2]], 'cost': r[3], 'race': r[4], 'atk': r[5],
                  'fx': [x for x in (norm(r[6]), norm(r[7])) if x]})
leaders = [{'id': r[0], 'name': r[1], 'race': r[2], 'hp': r[3]} for r in wb['主角圖鑑'].iter_rows(min_row=2, values_only=True) if r[0]]
(root / 'data' / 'cards.json').write_text(json.dumps({'cards': cards, 'leaders': leaders}, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
print(len(cards), 'cards,', len(leaders), 'leaders')
